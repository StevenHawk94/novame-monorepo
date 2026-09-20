#!/usr/bin/env python3
"""Compile the reviewed Connection Insight workbook into runtime JSON."""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


def pool(value):
    out = []
    pattern = re.compile(r"(.+?)\s*\((Positive|Negative),\s*(WYMM|TWL)\)")
    for raw in str(value or "").split(";"):
        raw = raw.strip()
        if not raw:
            continue
        match = pattern.fullmatch(raw)
        if not match:
            raise ValueError(f"Invalid Between You pool entry: {raw}")
        out.append({"group": match.group(1).strip(), "emotion": match.group(2), "source": match.group(3)})
    return out


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: build-connection-v8.py INPUT.xlsx OUTPUT.json")
    source, output = map(Path, sys.argv[1:])
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = [row for row in workbook["Connection Insight"].iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]
    templates = []
    for index, row in enumerate(rows, 2):
        category, group, sub, emotion, tag, title, description, ways_tag, ways_title, ways_description, ways_action = row
        templates.append({
            "sourceRow": index,
            "category": category,
            "group": group,
            "subScenario": sub,
            "emotion": emotion,
            "tag": tag,
            "title": title,
            "description": description,
            "waysIn": ({
                "tag": ways_tag,
                "title": ways_title,
                "description": ways_description,
                "action": ways_action,
            } if ways_tag or ways_title or ways_action else None),
        })

    mappings = []
    for index, row in enumerate(workbook["Between You Tag Pool Mapping"].iter_rows(min_row=2, values_only=True), 2):
        if not any(v is not None for v in row):
            continue
        tag, mapping_type, pool_a, pool_b, trigger = row
        parsed_a = pool(pool_a)
        parsed_b = parsed_a if str(pool_b or "").strip() == "(same as Pool A)" else pool(pool_b)
        mappings.append({
            "sourceRow": index,
            "tag": tag,
            "type": mapping_type,
            "poolA": parsed_a,
            "poolB": parsed_b,
            "triggerRule": trigger,
        })

    valid = {(row["group"], row["emotion"], row["category"]) for row in templates}
    invalid = [entry for mapping in mappings for entry in mapping["poolA"] + mapping["poolB"]
               if (entry["group"], entry["emotion"], entry["source"]) not in valid]
    if invalid:
        raise ValueError(f"Unknown Between You references: {invalid}")
    keys = [(row["category"], row["group"], row["subScenario"], row["emotion"]) for row in templates]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate Connection template key")

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({
        "version": "v8",
        "templates": templates,
        "betweenMappings": mappings,
    }, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"wrote {len(templates)} templates and {len(mappings)} Between You mappings to {output}")


if __name__ == "__main__":
    main()
