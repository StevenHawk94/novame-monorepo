#!/usr/bin/env python3
"""Build the reviewed Connection v6 SQL seed from its source workbook."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl


SECTION = {
    "WHAT YOU MAY HAVE MISSED": "missed",
    "THEIR WORLD LATELY": "world",
    "WAYS IN": "ways_in",
    "BETWEEN YOU LATELY": "between",
}


def key(value: str) -> str:
    return re.sub(r"^_+|_+$", "", re.sub(r"[^a-z0-9]+", "_", value.lower()))


def rows(workbook, sheet_name: str) -> list[dict]:
    sheet = workbook[sheet_name]
    headers = [sheet.cell(5, column).value for column in range(1, sheet.max_column + 1)]
    return [
        dict(zip(headers, [sheet.cell(row, column).value for column in range(1, sheet.max_column + 1)]))
        for row in range(6, sheet.max_row + 1)
        if sheet.cell(row, 1).value
    ]


def module_key(section: str, family_key: str) -> str | None:
    if section == "missed":
        return "worth_knowing"
    if section == "world":
        return "what_theyre_into" if family_key == "developing_interest" else "recent_vibe"
    if section == "between":
        return "shared_rhythm"
    return None


def compact_json(value) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def build(workbook_path: Path) -> str:
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    scenarios = rows(workbook, "Scenario Keys")
    variants = rows(workbook, "Template Variants")
    variants_by_key = {}
    for variant in variants:
        variants_by_key.setdefault(variant["Scenario Key"], []).append(variant)

    assert len(scenarios) == 220
    assert len(variants) == 880
    assert all(len(variants_by_key.get(row["Scenario Key"], [])) == 4 for row in scenarios)

    scenario_records = []
    for row in scenarios:
        scenario_key = row["Scenario Key"]
        family_key = key(row["Scenario Family"])
        section = SECTION[row["Section"]]
        lower_warm = next(
            variant for variant in variants_by_key[scenario_key]
            if variant["Depth Band"] == "lower" and variant["Tone Mode"] == "warm_clear"
        )
        scenario_records.append({
            "template_id": row["Legacy ID"],
            "family_key": family_key,
            "section": section,
            "module_key": module_key(section, family_key),
            "scenario_key": scenario_key,
            "scenario": row["Scenario"],
            "required_evidence": row["Required Evidence"],
            "disqualifiers": row["Disqualifiers"],
            "emotional_weight": row["Emotional Weight"],
            "tone_mode": "warm_clear",
            "depth_level": lower_warm["Depth"],
            "field_pattern": lower_warm["Field Pattern"],
            "label_options": row["Search Aliases"],
            "template_card": {
                "label": lower_warm["Label"],
                "title": lower_warm["Title"],
                "observation": lower_warm["Observation Template"],
                "meaning": lower_warm["Meaning Template"],
                "takeaway": lower_warm["Takeaway Template"],
            },
            "blocked_overlap": row["Blocked Overlap"],
            "signal_type": row["Signal Type"],
            "temporal_state": row["Temporal State"],
            "persistence": row["Persistence"],
            "topic_domain": row["Topic Domain"],
            "emotion_family": row["Emotion Family"],
            "support_mode": row["Support Mode"],
            "support_openness": row["Support Openness"],
            "response_preference": row["Response Preference"],
            "boundary_key": row["Boundary"],
            "timing": row["Timing"],
            "mutuality": row["Mutuality"],
            "depth_range": row["Depth Range"],
            "hard_match_tags": row["Hard Match Tags"],
            "preferred_tags": row["Preferred Tags"],
            "search_aliases": row["Search Aliases"],
            "retrieval_text": row["Retrieval Text"],
            "fallback_policy": row["Fallback Policy"],
            "library_version": "v6",
        })

    variant_records = [{
        "template_variant_id": row["Template Variant ID"],
        "scenario_key": row["Scenario Key"],
        "section": SECTION[row["Section"]],
        "variant_key": row["Variant"],
        "depth_band": row["Depth Band"],
        "tone_mode": row["Tone Mode"],
        "playful_eligible": bool(row["Playful Eligible"]),
        "playful_block_reason": row["Playful Block Reason"],
        "emotional_weight": row["Emotional Weight"],
        "depth_level": row["Depth"],
        "field_pattern": row["Field Pattern"],
        "template_card": {
            "label": row["Label"],
            "title": row["Title"],
            "observation": row["Observation Template"],
            "meaning": row["Meaning Template"],
            "takeaway": row["Takeaway Template"],
        },
        "required_slots": row["Required Slots"],
        "selection_rule": row["Selection Rule"],
        "null_policy": row["Null Policy"],
        "library_version": "v6",
    } for row in variants]

    return f"""-- Connection Insight v6: deterministic Scenario matching and 880 reviewed card variants.
-- Generated from Burrow_Connection_Insight_System_v6.xlsx.

begin;

alter table public.connection_scenario_templates
  add column if not exists signal_type text,
  add column if not exists temporal_state text,
  add column if not exists persistence text,
  add column if not exists topic_domain text,
  add column if not exists emotion_family text,
  add column if not exists support_mode text,
  add column if not exists support_openness text,
  add column if not exists response_preference text,
  add column if not exists boundary_key text,
  add column if not exists timing text,
  add column if not exists mutuality text,
  add column if not exists depth_range text,
  add column if not exists hard_match_tags text,
  add column if not exists preferred_tags text,
  add column if not exists search_aliases text,
  add column if not exists retrieval_text text,
  add column if not exists fallback_policy text;

create index if not exists connection_scenario_templates_runtime_match
  on public.connection_scenario_templates(family_key, section, signal_type, active);

create table if not exists public.connection_template_variants (
  template_variant_id text primary key,
  scenario_key text not null references public.connection_scenario_templates(scenario_key)
    on update cascade on delete cascade,
  section text not null check (section in ('missed','world','ways_in','between')),
  variant_key text not null,
  depth_band text not null check (depth_band in ('lower','upper')),
  tone_mode text not null check (tone_mode in ('warm_clear','playful_close')),
  playful_eligible boolean not null default false,
  playful_block_reason text,
  emotional_weight text,
  depth_level text not null,
  field_pattern text not null,
  template_card jsonb not null check (jsonb_typeof(template_card) = 'object'),
  required_slots text,
  selection_rule text,
  null_policy text,
  review_status text not null default 'approved' check (review_status in ('approved','disabled')),
  library_version text not null default 'v6',
  active boolean not null default true,
  created_at timestamptz not null default now(),
  updated_at timestamptz not null default now(),
  unique (scenario_key, depth_band, tone_mode)
);
create index if not exists connection_template_variants_runtime
  on public.connection_template_variants(scenario_key, depth_band, tone_mode, active);

alter table public.connection_template_variants enable row level security;
drop policy if exists connection_template_variants_service on public.connection_template_variants;
create policy connection_template_variants_service on public.connection_template_variants
  for all to service_role using (true) with check (true);
revoke all on public.connection_template_variants from public, anon, authenticated;
grant all on public.connection_template_variants to service_role;

update public.connection_scenario_templates set active = false
where library_version <> 'v6';

insert into public.connection_scenario_templates(
  template_id,family_key,section,module_key,scenario_key,scenario,required_evidence,
  disqualifiers,emotional_weight,tone_mode,depth_level,field_pattern,label_options,
  template_card,blocked_overlap,signal_type,temporal_state,persistence,topic_domain,
  emotion_family,support_mode,support_openness,response_preference,boundary_key,timing,
  mutuality,depth_range,hard_match_tags,preferred_tags,search_aliases,retrieval_text,
  fallback_policy,review_status,library_version,active
)
select template_id,family_key,section,module_key,scenario_key,scenario,required_evidence,
  disqualifiers,emotional_weight,tone_mode,depth_level,field_pattern,label_options,
  template_card,blocked_overlap,signal_type,temporal_state,persistence,topic_domain,
  emotion_family,support_mode,support_openness,response_preference,boundary_key,timing,
  mutuality,depth_range,hard_match_tags,preferred_tags,search_aliases,retrieval_text,
  fallback_policy,'approved','v6',true
from jsonb_to_recordset($connection_scenarios_v6$
{compact_json(scenario_records)}
$connection_scenarios_v6$::jsonb) as x(
  template_id text,family_key text,section text,module_key text,scenario_key text,
  scenario text,required_evidence text,disqualifiers text,emotional_weight text,
  tone_mode text,depth_level text,field_pattern text,label_options text,template_card jsonb,
  blocked_overlap text,signal_type text,temporal_state text,persistence text,
  topic_domain text,emotion_family text,support_mode text,support_openness text,
  response_preference text,boundary_key text,timing text,mutuality text,depth_range text,
  hard_match_tags text,preferred_tags text,search_aliases text,retrieval_text text,
  fallback_policy text
)
on conflict (template_id) do update set
  family_key=excluded.family_key,section=excluded.section,module_key=excluded.module_key,
  scenario_key=excluded.scenario_key,scenario=excluded.scenario,
  required_evidence=excluded.required_evidence,disqualifiers=excluded.disqualifiers,
  emotional_weight=excluded.emotional_weight,tone_mode=excluded.tone_mode,
  depth_level=excluded.depth_level,field_pattern=excluded.field_pattern,
  label_options=excluded.label_options,template_card=excluded.template_card,
  blocked_overlap=excluded.blocked_overlap,signal_type=excluded.signal_type,
  temporal_state=excluded.temporal_state,persistence=excluded.persistence,
  topic_domain=excluded.topic_domain,emotion_family=excluded.emotion_family,
  support_mode=excluded.support_mode,support_openness=excluded.support_openness,
  response_preference=excluded.response_preference,boundary_key=excluded.boundary_key,
  timing=excluded.timing,mutuality=excluded.mutuality,depth_range=excluded.depth_range,
  hard_match_tags=excluded.hard_match_tags,preferred_tags=excluded.preferred_tags,
  search_aliases=excluded.search_aliases,retrieval_text=excluded.retrieval_text,
  fallback_policy=excluded.fallback_policy,review_status='approved',library_version='v6',
  active=true,updated_at=now();

update public.connection_template_variants set active = false
where library_version <> 'v6';

insert into public.connection_template_variants(
  template_variant_id,scenario_key,section,variant_key,depth_band,tone_mode,
  playful_eligible,playful_block_reason,emotional_weight,depth_level,field_pattern,
  template_card,required_slots,selection_rule,null_policy,review_status,library_version,active
)
select template_variant_id,scenario_key,section,variant_key,depth_band,tone_mode,
  playful_eligible,playful_block_reason,emotional_weight,depth_level,field_pattern,
  template_card,required_slots,selection_rule,null_policy,'approved','v6',true
from jsonb_to_recordset($connection_variants_v6$
{compact_json(variant_records)}
$connection_variants_v6$::jsonb) as x(
  template_variant_id text,scenario_key text,section text,variant_key text,
  depth_band text,tone_mode text,playful_eligible boolean,playful_block_reason text,
  emotional_weight text,depth_level text,field_pattern text,template_card jsonb,
  required_slots text,selection_rule text,null_policy text
)
on conflict (template_variant_id) do update set
  scenario_key=excluded.scenario_key,section=excluded.section,variant_key=excluded.variant_key,
  depth_band=excluded.depth_band,tone_mode=excluded.tone_mode,
  playful_eligible=excluded.playful_eligible,
  playful_block_reason=excluded.playful_block_reason,
  emotional_weight=excluded.emotional_weight,depth_level=excluded.depth_level,
  field_pattern=excluded.field_pattern,template_card=excluded.template_card,
  required_slots=excluded.required_slots,selection_rule=excluded.selection_rule,
  null_policy=excluded.null_policy,review_status='approved',library_version='v6',
  active=true,updated_at=now();

commit;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.write_text(build(args.workbook), encoding="utf-8")


if __name__ == "__main__":
    main()
